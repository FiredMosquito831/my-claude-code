"""The Exit column and the Exit filter over HTTP (7.88.0).

The list route adds one key per row, ``exit``; every route the Analytics page
calls takes the ``exit`` filter and counts the same rows; the export gains the
two Exit columns only when its ``exit`` field is asked for, and keeps Folder.
Every address is a documentation address (RFC 5737) or a made-up name.
"""

import csv
import io
import json
import time

import pytest
from fastapi.testclient import TestClient

from my_claude_code.core.request_log import (
    RequestRecord,
    RouteAttempt,
    RouteAttemptOutcome,
    get_request_log_store,
)
from tests.api.support import create_test_app

DEMO = "C:\\Users\\devuser\\Projects\\demo"


def _dials(*labels: str) -> dict:
    return {
        "ladder": {
            "tries": [{"source": "upstream", "status": 200, "proxy": labels[-1]}],
            "summary": {"tries": 1},
            "credentials": [],
            "dials": [{"at_try": 0, "proxy": label} for label in labels],
        }
    }


@pytest.fixture
def client():
    return TestClient(create_test_app(), client=("127.0.0.1", 50000))


@pytest.fixture
def seeded_store(tmp_path):
    store = get_request_log_store(tmp_path / "requests.db")
    assert store is not None
    base = time.time() - 60
    plan = [
        ("chain", "Tokyo exit", None),
        ("direct", "direct", None),
        (
            "rotated",
            "exit-b.example:1080",
            _dials("exit-a.example:1080", "exit-b.example:1080"),
        ),
        ("bare", None, None),
    ]
    for index, (request_id, label, params) in enumerate(plan):
        store.enqueue(
            RequestRecord(
                id=request_id,
                endpoint="/v1/messages",
                protocol="anthropic",
                provider="p1",
                resolved_model="m1",
                harness="claude",
                ts_epoch=base + index,
                status="success",
                tokens_in=10,
                tokens_out=2,
                duration_ms=100.0,
                ttft_ms=20.0,
                input_text="hi",
                output_text="out",
                project_dir=DEMO,
                route_attempt=0,
                attempts=(
                    RouteAttempt(
                        attempt=0,
                        provider="p1",
                        model_ref="p1/m1",
                        outcome=RouteAttemptOutcome.SUCCEEDED,
                        params=params,
                        proxy_label=label,
                    ),
                ),
            )
        )
    store.close()
    yield store


def test_the_list_route_adds_the_exit_of_each_row(client, seeded_store) -> None:
    listed = client.get("/admin/api/requests", params={"limit": 50}).json()
    exits = {row["id"]: row["exit"] for row in listed["rows"]}

    assert exits == {
        "chain": {"label": "Tokyo exit", "tried": ["Tokyo exit"]},
        "direct": {"label": "direct", "tried": ["direct"]},
        "rotated": {
            "label": "exit-b.example:1080",
            "tried": ["exit-a.example:1080", "exit-b.example:1080"],
        },
        "bare": {"label": None, "tried": []},
    }
    # Folder is still on every row the table reads, for the chip and the title.
    assert {row["project_dir"] for row in listed["rows"]} == {DEMO}


@pytest.mark.parametrize(
    ("params", "expected"),
    [
        ({"exit": "tokyo"}, 1),
        ({"exit": "exit-a"}, 1),
        ({"exit": "1080"}, 1),
        ({"exit": "direct"}, 1),
        ({"exit": "nowhere"}, 0),
        ({"exit": "  "}, 4),
        ({"exit": "tokyo", "folder": "demo"}, 1),
        ({}, 4),
    ],
)
def test_every_analytics_route_honours_the_exit_filter(
    client, seeded_store, params, expected
) -> None:
    listed = client.get("/admin/api/requests", params={"limit": 50, **params}).json()
    assert listed["total"] == expected
    assert client.get("/admin/api/requests/count", params=params).json()["total"] == (
        expected
    )
    assert client.get("/admin/api/requests/stats", params=params).json()["total"] == (
        expected
    )
    assert client.get("/admin/api/requests/pulse", params=params).json()["total"] == (
        expected
    )
    assert client.get("/admin/api/requests/ttft", params=params).json()["measured"] == (
        expected
    )
    cost = client.get("/admin/api/requests/cost", params=params).json()
    assert cost["totals"]["requests"] == expected
    origin = client.get("/admin/api/requests/origin", params=params).json()
    assert sum(row["requests"] for row in origin["by_folder"]) == expected
    no_answer = client.get("/admin/api/requests/no-answer", params=params)
    assert no_answer.status_code == 200


def _csv(client, **params: str) -> list[list[str]]:
    response = client.get(
        "/admin/api/export", params={"format": "csv", "scope": "requests", **params}
    )
    assert response.status_code == 200, response.text
    return list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))


def test_the_export_gains_the_exit_columns_only_when_asked(
    client, seeded_store
) -> None:
    default = _csv(client)
    assert "Exit" not in default[0]
    assert "Exits tried" not in default[0]

    asked = _csv(client, fields="providers,origin,exit")
    header = asked[0]
    assert "Folder" in header
    rows = {row[header.index("ID")]: row for row in asked[1:]}
    exit_col, exits_col = header.index("Exit"), header.index("Exits tried")
    assert rows["rotated"][exit_col] == "exit-b.example:1080"
    assert rows["rotated"][exits_col] == "exit-a.example:1080; exit-b.example:1080"
    assert rows["bare"][exit_col] == ""
    assert rows["chain"][header.index("Folder")] == DEMO


def test_the_xlsx_export_carries_the_exit_columns(client, seeded_store) -> None:
    pytest.importorskip("openpyxl")
    import openpyxl

    response = client.get(
        "/admin/api/export",
        params={"format": "xlsx", "scope": "requests", "fields": "origin,exit"},
    )
    sheet = openpyxl.load_workbook(io.BytesIO(response.content)).active
    header = [sheet.cell(1, col).value for col in range(1, sheet.max_column + 1)]
    assert {"Folder", "Exit", "Exits tried"} <= set(header)
    exits = {
        sheet.cell(row, header.index("ID") + 1).value: sheet.cell(
            row, header.index("Exit") + 1
        ).value
        for row in range(2, sheet.max_row + 1)
    }
    assert exits["chain"] == "Tokyo exit"
    assert exits["direct"] == "direct"


@pytest.mark.parametrize("scope", ["requests", "attempts"])
def test_the_export_follows_the_exit_filter(client, seeded_store, scope) -> None:
    response = client.get(
        "/admin/api/export",
        params={"format": "json", "scope": scope, "exit": "exit-a"},
    )
    assert response.status_code == 200, response.text
    payload = json.loads(response.content)
    rows = payload["rows"] if isinstance(payload, dict) else payload
    key = "id" if scope == "requests" else "request_id"
    assert [row[key] for row in rows] == ["rotated"]
