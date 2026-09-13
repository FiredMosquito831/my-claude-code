"""The attempt-level export: one row per attempt, not per request.

Until 7.6.0 no export path read ``request_attempts`` except the per-request
ladder rollup, so the models a fallback chain tried and abandoned -- the ones
that spend the time -- were unexportable. These tests pin the row shape, the
four formats, the filter/scope reuse, and the rule that a number nobody
measured comes out empty rather than zero.
"""

import csv
import io
import json
import time

import pytest
from fastapi.testclient import TestClient

from my_claude_code.core import export as export_engine
from my_claude_code.core.request_log import (
    RequestRecord,
    RouteAttempt,
    RouteAttemptOutcome,
    get_request_log_store,
)
from tests.api.support import create_test_app


@pytest.fixture
def client():
    return TestClient(create_test_app(), client=("127.0.0.1", 50000))


@pytest.fixture
def chain_store(tmp_path):
    """One fallback chain, one clean request, one row with no attempts at all."""
    store = get_request_log_store(tmp_path / "requests.db")
    assert store is not None
    base = time.time() - 100
    store.enqueue(
        RequestRecord(
            id="chain",
            endpoint="/v1/messages",
            protocol="anthropic",
            provider="p1",
            requested_model="mcc/best",
            resolved_model="p1/fast",
            harness="claude_code",
            ts_epoch=base + 10,
            status="success",
            duration_ms=5000.0,
            ttft_ms=4712.0,
            ttft_winner_ms=300.0,
            attempts=(
                RouteAttempt(
                    attempt=0,
                    provider="p1",
                    model_ref="p1/slow",
                    outcome=RouteAttemptOutcome.FAILED,
                    error_kind="rate_limit",
                    error_message="429 from the host",
                    duration_ms=4400.0,
                    ttft_ms=None,
                    key_label="key-a",
                    key_index=0,
                    ladder_tries=3,
                    params={
                        "ladder": {"root_cause": "429 rate_limit"},
                        "wire": {"surface": "responses"},
                        "early_retries": 2,
                    },
                ),
                RouteAttempt(
                    attempt=1,
                    provider="p1",
                    model_ref="p1/benched",
                    outcome=RouteAttemptOutcome.SKIPPED,
                    error_message="benched: 3 consecutive failures, 22 s left",
                    params={
                        "bench": {
                            "mode": "consecutive",
                            "failures": 3,
                            "last_kind": "server_error",
                            "last_status": 500,
                            "remaining_seconds": 22.0,
                        }
                    },
                ),
                RouteAttempt(
                    attempt=2,
                    provider="p1",
                    model_ref="p1/fast",
                    outcome=RouteAttemptOutcome.SUCCEEDED,
                    duration_ms=600.0,
                    ttft_ms=300.0,
                    first_reasoning_ms=120.0,
                    tokens_out=64,
                    key_label="key-b",
                    key_index=1,
                    ladder_tries=1,
                ),
            ),
        )
    )
    store.enqueue(
        RequestRecord(
            id="clean",
            endpoint="/v1/responses",
            protocol="openai",
            provider="p2",
            requested_model="p2/m",
            resolved_model="p2/m",
            harness="codex",
            ts_epoch=base + 20,
            status="success",
            duration_ms=900.0,
            attempts=(
                RouteAttempt(
                    attempt=0,
                    provider="p2",
                    model_ref="p2/m",
                    outcome=RouteAttemptOutcome.SUCCEEDED,
                ),
            ),
        )
    )
    # No attempts recorded at all: every row written before the side table
    # existed looks like this, and it must contribute nothing rather than one
    # row of blanks pretending an attempt happened.
    store.enqueue(
        RequestRecord(
            id="attemptless",
            endpoint="/v1/messages",
            protocol="anthropic",
            provider="p1",
            resolved_model="p1/fast",
            harness="claude_code",
            ts_epoch=base + 30,
            status="success",
        )
    )
    store.close()
    yield store


def _export(client, **params):
    return client.get("/admin/api/export", params=params)


class TestRowShape:
    def test_one_row_per_attempt_newest_request_first(
        self, client, chain_store
    ) -> None:
        rows = _export(client, format="json", scope="attempts").json()

        # Newest request first, attempts in chain order within it, so a
        # fallback chain reads top to bottom.
        assert [(row["request_id"], row["attempt"]) for row in rows] == [
            ("clean", 0),
            ("chain", 0),
            ("chain", 1),
            ("chain", 2),
        ]

    def test_a_request_with_no_attempts_contributes_no_rows(
        self, client, chain_store
    ) -> None:
        rows = _export(client, format="json", scope="attempts").json()

        assert all(row["request_id"] != "attemptless" for row in rows)

    def test_the_parent_requests_dimensions_ride_on_every_attempt(
        self, client, chain_store
    ) -> None:
        rows = [
            row
            for row in _export(client, format="json", scope="attempts").json()
            if row["request_id"] == "chain"
        ]

        # A failed attempt has no endpoint, harness or requested model of its
        # own; those are facts about the request it belongs to.
        assert {row["harness"] for row in rows} == {"claude_code"}
        assert {row["endpoint"] for row in rows} == {"/v1/messages"}
        assert {row["requested_model"] for row in rows} == {"mcc/best"}
        assert {row["resolved_model"] for row in rows} == {"p1/fast"}
        assert {row["request_status"] for row in rows} == {"success"}
        assert {row["ts_iso"] for row in rows} == {rows[0]["ts_iso"]}

    def test_each_attempt_carries_its_own_model_and_verdict(
        self, client, chain_store
    ) -> None:
        rows = {
            (row["request_id"], row["attempt"]): row
            for row in _export(client, format="json", scope="attempts").json()
        }

        assert rows[("chain", 0)]["attempt_model"] == "p1/slow"
        assert rows[("chain", 0)]["outcome"] == "failed"
        assert rows[("chain", 1)]["attempt_model"] == "p1/benched"
        assert rows[("chain", 1)]["outcome"] == "skipped"
        assert rows[("chain", 2)]["attempt_model"] == "p1/fast"
        assert rows[("chain", 2)]["outcome"] == "succeeded"

    def test_the_winners_latency_is_its_own_not_the_chains(
        self, client, chain_store
    ) -> None:
        rows = {
            (row["request_id"], row["attempt"]): row
            for row in _export(client, format="json", scope="attempts").json()
        }

        winner = rows[("chain", 2)]
        assert winner["ttft_ms"] == 300.0
        assert winner["first_reasoning_ms"] == 120.0
        assert winner["duration_ms"] == 600.0
        assert winner["tokens_out"] == 64

    def test_an_unmeasured_number_is_empty_never_zero(
        self, client, chain_store
    ) -> None:
        rows = {
            (row["request_id"], row["attempt"]): row
            for row in _export(client, format="json", scope="attempts").json()
        }

        loser = rows[("chain", 0)]
        # It never produced answer content, so it has no TTFT. Zero would read
        # as "instant".
        assert loser["ttft_ms"] is None
        assert loser["first_reasoning_ms"] is None
        # ``tokens_out`` is filled only on the winning attempt; the request's
        # count belongs to the request, not to a model that did not answer.
        assert loser["tokens_out"] is None
        assert loser["cost_usd"] is None

        skipped = rows[("chain", 1)]
        assert skipped["duration_ms"] is None
        assert skipped["ttft_ms"] is None


class TestDerivedColumns:
    def test_ended_by_says_what_stopped_each_attempt(self, client, chain_store) -> None:
        rows = {
            (row["request_id"], row["attempt"]): row
            for row in _export(client, format="json", scope="attempts").json()
        }

        assert rows[("chain", 0)]["ended_by"] == "rate_limit"
        assert rows[("chain", 1)]["ended_by"] == "benched (server_error)"
        assert rows[("chain", 2)]["ended_by"] == "succeeded"

    def test_the_skip_reason_survives_as_both_prose_and_structure(
        self, client, chain_store
    ) -> None:
        rows = {
            (row["request_id"], row["attempt"]): row
            for row in _export(client, format="json", scope="attempts").json()
        }

        skipped = rows[("chain", 1)]
        assert "benched" in skipped["error_message"]
        assert skipped["bench_reason"] == (
            "consecutive · 3 failures · last 500 server_error · 22 s left"
        )
        # An attempt that actually ran was not benched, and says so by
        # omission rather than by an invented value.
        assert rows[("chain", 2)]["bench_reason"] is None

    def test_the_ladder_root_cause_is_the_attempts_own(
        self, client, chain_store
    ) -> None:
        rows = {
            (row["request_id"], row["attempt"]): row
            for row in _export(client, format="json", scope="attempts").json()
        }

        assert rows[("chain", 0)]["ladder_tries"] == 3
        assert rows[("chain", 0)]["ladder_root_cause"] == "429 rate_limit"
        assert rows[("chain", 2)]["ladder_root_cause"] is None

    def test_optional_groups_are_absent_until_selected(
        self, client, chain_store
    ) -> None:
        default = _export(client, format="json", scope="attempts").json()[0]
        assert "wire_surface" not in default
        assert "early_retries" not in default

        picked = _export(
            client, format="json", scope="attempts", fields="wire,recovery"
        ).json()
        row = next(
            row
            for row in picked
            if row["attempt"] == 0 and row["request_id"] == "chain"
        )
        assert row["wire_surface"] == "responses"
        assert row["early_retries"] == 2
        # Never counted for this attempt, so absent -- not zero.
        assert row["salvages"] is None
        assert row["key_index"] == 0

    def test_the_private_params_blob_is_never_an_output_column(
        self, client, chain_store
    ) -> None:
        rows = _export(
            client,
            format="json",
            scope="attempts",
            fields="failure,ladder,wire,recovery",
        ).json()

        assert all(export_engine.ATTEMPT_PARAMS_KEY not in row for row in rows)


class TestFormats:
    @pytest.mark.parametrize(
        ("fmt", "content_type", "extension"),
        (
            ("json", "application/json", "json"),
            ("csv", "text/csv", "csv"),
            (
                "xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "xlsx",
            ),
            ("txt", "text/plain", "txt"),
        ),
    )
    def test_every_format_streams_with_the_attempt_filename(
        self, client, chain_store, fmt, content_type, extension
    ) -> None:
        response = _export(client, format=fmt, scope="attempts")

        assert response.status_code == 200
        assert response.headers["content-type"].startswith(content_type)
        disposition = response.headers["content-disposition"]
        assert disposition.startswith("attachment")
        assert "mcc-attempts-" in disposition
        assert f".{extension}" in disposition

    def test_csv_labels_the_columns_and_leaves_unmeasured_cells_empty(
        self, client, chain_store
    ) -> None:
        response = _export(client, format="csv", scope="attempts")
        reader = csv.reader(io.StringIO(response.content.decode("utf-8-sig")))
        header = next(reader)
        body = list(reader)

        assert header[:3] == ["Request ID", "Time", "Attempt #"]
        for label in ("Attempt model", "Outcome", "Ended by", "TTFT (ms)"):
            assert label in header
        ttft = header.index("TTFT (ms)")
        loser = next(row for row in body if row[0] == "chain" and row[2] == "0")
        # CSV renders a missing measurement as an empty cell, exactly as the
        # request export does -- not as "0" and not as "None".
        assert loser[ttft] == ""

    def test_json_renders_a_missing_measurement_as_null(
        self, client, chain_store
    ) -> None:
        payload = json.loads(_export(client, format="json", scope="attempts").content)
        loser = next(
            row
            for row in payload
            if row["request_id"] == "chain" and row["attempt"] == 0
        )

        assert loser["ttft_ms"] is None

    def test_xlsx_reads_back_with_one_row_per_attempt(
        self, client, chain_store
    ) -> None:
        pytest.importorskip("openpyxl")
        import io as _io

        import openpyxl

        response = _export(client, format="xlsx", scope="attempts")
        sheet = openpyxl.load_workbook(_io.BytesIO(response.content)).active

        assert sheet.max_row == 5  # header + 4 attempts
        header = [sheet.cell(1, col).value for col in range(1, sheet.max_column + 1)]
        assert "Attempt model" in header
        assert "Ended by" in header

    def test_txt_is_readable(self, client, chain_store) -> None:
        text = _export(client, format="txt", scope="attempts").content.decode("utf-8")

        assert "Attempt model" in text
        assert "p1/slow" in text


class TestFiltersAndScope:
    def test_the_time_window_is_the_parent_requests(self, client, chain_store) -> None:
        everything = _export(client, format="json", scope="attempts").json()
        recent = _export(
            client,
            format="json",
            scope="attempts",
            since=str(time.time() - 85),
        ).json()

        assert len(everything) == 4
        # Only ``clean`` (base+20) and ``attemptless`` (base+30) fall inside,
        # and the second has no attempts.
        assert {row["request_id"] for row in recent} == {"clean"}

    def test_the_harness_filter_selects_whole_chains(self, client, chain_store) -> None:
        rows = _export(client, format="json", scope="attempts", harness="codex").json()

        assert {row["request_id"] for row in rows} == {"clean"}

    def test_the_local_filter_is_validated_like_the_request_export(
        self, client, chain_store
    ) -> None:
        assert (
            _export(
                client, format="json", scope="attempts", local="sometimes"
            ).status_code
            == 422
        )
        assert (
            _export(client, format="json", scope="attempts", local="all").status_code
            == 200
        )

    def test_an_invalid_status_is_refused(self, client, chain_store) -> None:
        response = _export(client, format="json", scope="attempts", status="nope")

        assert response.status_code == 422

    def test_a_model_filter_matches_the_request_not_the_attempt(
        self, client, chain_store
    ) -> None:
        """The filters are the request-log filters, deliberately.

        ``model=p1/fast`` selects the request whose resolved model is that, and
        therefore *every* attempt it made -- including the two that failed.
        That is the question the export exists to answer; filtering to the
        matching attempt alone would hide the chain it belonged to.
        """
        rows = _export(client, format="json", scope="attempts", model="p1/fast").json()

        assert {(row["request_id"], row["attempt"]) for row in rows} == {
            ("chain", 0),
            ("chain", 1),
            ("chain", 2),
        }

    def test_an_unknown_field_is_a_400(self, client, chain_store) -> None:
        response = _export(client, format="json", scope="attempts", fields="nope")

        assert response.status_code == 400

    def test_grouping_is_refused_rather_than_ignored(self, client, chain_store) -> None:
        """Detail-only, and it says so.

        Averaging latency across attempts of different models inside one
        request is the number the per-model cards already answer properly. A
        silently dropped ``group_by`` would hand back a flat file the caller
        believes is grouped.
        """
        response = _export(client, format="json", scope="attempts", group_by="provider")

        assert response.status_code == 400
        assert "cannot be grouped" in response.json()["detail"]

    def test_a_remote_client_is_forbidden(self, client, chain_store) -> None:
        remote = TestClient(client.app, client=("203.0.113.10", 50000))

        response = remote.get(
            "/admin/api/export", params={"format": "json", "scope": "attempts"}
        )

        assert response.status_code == 403


class TestRedaction:
    """Redaction is a write-time property, and the attempt export inherits it.

    ``request_attempts.error_message`` is written through
    ``safe_exception_message`` -> ``redact_sensitive_error_text``
    (``application/execution.py``), so a credential-shaped string never reaches
    the column the export reads. This pins that the export does not somehow
    route around it.
    """

    def test_a_credential_shaped_error_never_reaches_the_download(
        self, client, tmp_path
    ) -> None:
        from my_claude_code.core.diagnostics import safe_exception_message

        secret = "AKIAIOSFODNN7EXAMPLE"
        stored = safe_exception_message(RuntimeError(f"denied for key {secret}"))
        assert secret not in stored

        store = get_request_log_store(tmp_path / "requests.db")
        assert store is not None
        store.enqueue(
            RequestRecord(
                id="leaky",
                endpoint="/v1/messages",
                protocol="anthropic",
                provider="p1",
                resolved_model="p1/m",
                ts_epoch=time.time() - 5,
                status="error",
                attempts=(
                    RouteAttempt(
                        attempt=0,
                        provider="p1",
                        model_ref="p1/m",
                        outcome=RouteAttemptOutcome.FAILED,
                        error_kind="auth",
                        error_message=stored,
                    ),
                ),
            )
        )
        store.close()

        for fmt in ("json", "csv", "txt"):
            body = _export(client, format=fmt, scope="attempts").content
            assert secret.encode() not in body
            assert b"<redacted>" in body


class TestStreamingIsBounded:
    def test_more_attempts_than_a_page_all_arrive(self, client, tmp_path) -> None:
        """The walk pages over ``requests``; nothing is capped at one page."""
        store = get_request_log_store(tmp_path / "requests.db")
        assert store is not None
        base = time.time() - 10_000
        for index in range(600):
            store.enqueue(
                RequestRecord(
                    id=f"r{index}",
                    endpoint="/v1/messages",
                    protocol="anthropic",
                    provider="p1",
                    resolved_model="p1/m",
                    ts_epoch=base + index,
                    status="success",
                    attempts=(
                        RouteAttempt(
                            attempt=0,
                            provider="p1",
                            model_ref="p1/a",
                            outcome=RouteAttemptOutcome.FAILED,
                        ),
                        RouteAttempt(
                            attempt=1,
                            provider="p1",
                            model_ref="p1/b",
                            outcome=RouteAttemptOutcome.SUCCEEDED,
                        ),
                    ),
                )
            )
        store.close()

        rows = _export(client, format="json", scope="attempts").json()

        assert len(rows) == 1_200

    def test_the_store_iterator_pages_rather_than_buffering(self, tmp_path) -> None:
        store = get_request_log_store(tmp_path / "requests.db")
        assert store is not None
        base = time.time() - 1_000
        for index in range(5):
            store.enqueue(
                RequestRecord(
                    id=f"p{index}",
                    endpoint="/v1/messages",
                    protocol="anthropic",
                    provider="p1",
                    resolved_model="p1/m",
                    ts_epoch=base + index,
                    status="success",
                    attempts=(
                        RouteAttempt(
                            attempt=0,
                            provider="p1",
                            model_ref="p1/a",
                            outcome=RouteAttemptOutcome.SUCCEEDED,
                        ),
                    ),
                )
            )
        store.close()

        # page_size 2 over 5 requests exercises the keyset cursor three times.
        rows = list(store.iter_export_attempt_rows(page_size=2))

        assert [row["request_id"] for row in rows] == ["p4", "p3", "p2", "p1", "p0"]


class TestEndedByIsNotGatedByTheFailureGroup:
    def test_it_still_names_the_error_when_failure_is_deselected(
        self, client, chain_store
    ) -> None:
        """``ended_by`` is always-derived, so its inputs must always be read.

        The store fetches ``error_kind`` whatever the field selection, and only
        the projection at the end drops it. Deselecting the failure group must
        therefore hide the column without emptying the summary derived from it.
        """
        rows = {
            (row["request_id"], row["attempt"]): row
            for row in _export(
                client, format="json", scope="attempts", fields="tokens"
            ).json()
        }

        assert "error_kind" not in rows[("chain", 0)]
        assert rows[("chain", 0)]["ended_by"] == "rate_limit"
        assert rows[("chain", 1)]["ended_by"] == "benched (server_error)"
